#!/usr/bin/env python3
# ==============================================================================
# Copyright (C) 2026 Intel Corporation
#
# SPDX-License-Identifier: MIT
# ==============================================================================
"""
Merge XLSX reports from re-runs of failed tests into the full (initial) XLSX report.

For every test case found in a retry report, the row in the full report is updated with
the retry result (last retry wins). Details of every attempt are kept in the 'details' sheet,
the 'Details' link points to the last attempt, and two columns are added to the 'main' sheet:
'Attempts' and 'Attempt history' (e.g. 'FAIL -> PASS'). A text summary (.txt) is regenerated
next to the output XLSX in the same format as produced by XlsxReporter.

Usage:
  merge_retry_reports.py --base <full.xlsx> --retries <retry1.xlsx> [<retry2.xlsx> ...] --output <merged.xlsx>
"""

import argparse
import sys
import warnings
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.hyperlink import Hyperlink

# XlsxReporter uses Excel 2010 data bars; openpyxl keeps them in basic form and warns about the extension
warnings.filterwarnings('ignore', message='.*extension is not supported.*', module='openpyxl')

# Layout produced by XlsxReporter (1-based columns/rows, as used by openpyxl)
MAIN_HEADER_ROW = 6
MAIN_FIRST_ROW = MAIN_HEADER_ROW + 1
COL_NUM, COL_SUITE, COL_TEST, COL_DURATION, COL_FPS, COL_RESULT = 1, 2, 3, 4, 5, 6
COL_FRAMES_TOTAL, COL_FRAMES_FAIL, COL_FRAMES_FAIL_PC, COL_FRAMES_MISSING = 7, 8, 9, 10
COL_OBJECTS_TOTAL, COL_OBJECTS_FAIL, COL_OBJECTS_FAIL_PC = 11, 12, 13
COL_DETAILS_LINK = 14
COL_ATTEMPTS = 15
COL_HISTORY = 16
COPIED_COLS = (COL_DURATION, COL_FPS, COL_RESULT, COL_FRAMES_TOTAL, COL_FRAMES_FAIL,
               COL_FRAMES_MISSING, COL_OBJECTS_TOTAL, COL_OBJECTS_FAIL)
# Columns holding per-row formulas: (column, numerator column, denominator column)
FORMULA_COLS = ((COL_FRAMES_FAIL_PC, COL_FRAMES_FAIL, COL_FRAMES_TOTAL),
                (COL_OBJECTS_FAIL_PC, COL_OBJECTS_FAIL, COL_OBJECTS_TOTAL))

DETAILS_HEADER_ROW = 1
DETAILS_COLS = 6  # NUM, RESULT, ERROR, PIPELINE, STDERR, STDOUT
COL_DETAILS_ATTEMPT = DETAILS_COLS + 1

FLAKY_FILL = PatternFill(start_color='FFC000', end_color='FFC000', fill_type='solid')


def status(value) -> str:
    return "PASS" if value is True else "FAIL"


def main_rows(sheet):
    for row in range(MAIN_FIRST_ROW, sheet.max_row + 1):
        if sheet.cell(row, COL_NUM).value is not None:
            yield row


def details_rows_by_num(sheet) -> dict:
    return {sheet.cell(row, 1).value: row
            for row in range(DETAILS_HEADER_ROW + 1, sheet.max_row + 1)
            if sheet.cell(row, 1).value is not None}


def copy_cell(src, dst):
    dst.value = src.value
    if src.has_style:
        dst.font = copy(src.font)
        dst.alignment = copy(src.alignment)
        dst.number_format = src.number_format


def check_layout(wb, path: Path):
    """Fail fast if report was not produced by XlsxReporter (e.g. XlsxBenchmarkReporter has other layout)"""
    expected = {COL_SUITE: "Test suite", COL_TEST: "Test", COL_RESULT: "Result"}
    if 'main' not in wb.sheetnames or 'details' not in wb.sheetnames or \
            any(wb['main'].cell(MAIN_HEADER_ROW, col).value != name for col, name in expected.items()):
        sys.exit(f"ERROR: {path} has unsupported layout (expected XlsxReporter functional tests report)")


def merge(base_path: Path, retry_paths: list, output_path: Path):
    wb = load_workbook(base_path)
    check_layout(wb, base_path)
    main, details = wb['main'], wb['details']

    # Index base rows by (suite, test); duplicated names are matched in order of appearance
    index = {}
    for row in main_rows(main):
        key = (main.cell(row, COL_SUITE).value, main.cell(row, COL_TEST).value)
        index.setdefault(key, []).append(row)
    history = {row: [status(main.cell(row, COL_RESULT).value)] for rows in index.values() for row in rows}

    for row in details_rows_by_num(details).values():
        details.cell(row, COL_DETAILS_ATTEMPT, 1)

    for attempt, retry_path in enumerate(retry_paths, start=2):
        rwb = load_workbook(retry_path)
        check_layout(rwb, retry_path)
        rmain, rdetails = rwb['main'], rwb['details']
        rdetails_by_num = details_rows_by_num(rdetails)
        seen = {}
        for rrow in main_rows(rmain):
            key = (rmain.cell(rrow, COL_SUITE).value, rmain.cell(rrow, COL_TEST).value)
            occurrence = seen.get(key, 0)
            seen[key] = occurrence + 1
            if key not in index or occurrence >= len(index[key]):
                print(f"WARNING: {key} from {retry_path.name} not found in base report - skipped", file=sys.stderr)
                continue
            row = index[key][occurrence]

            for col in COPIED_COLS:
                copy_cell(rmain.cell(rrow, col), main.cell(row, col))
            for col, num_col, den_col in FORMULA_COLS:
                copy_cell(rmain.cell(rrow, col), main.cell(row, col))
                if isinstance(main.cell(row, col).value, str) and main.cell(row, col).value.startswith('='):
                    num = main.cell(row, num_col).coordinate
                    den = main.cell(row, den_col).coordinate
                    main.cell(row, col).value = f'=IFERROR({num}/{den}, 0)'

            # Append details of this attempt and point the 'Details' link to it
            rdet_row = rdetails_by_num.get(rmain.cell(rrow, COL_NUM).value)
            if rdet_row is not None:
                det_row = details.max_row + 1
                for col in range(1, DETAILS_COLS + 1):
                    copy_cell(rdetails.cell(rdet_row, col), details.cell(det_row, col))
                details.cell(det_row, 1).value = main.cell(row, COL_NUM).value
                details.cell(det_row, COL_DETAILS_ATTEMPT, attempt)
                link = main.cell(row, COL_DETAILS_LINK)
                link.value = "Details"
                link.hyperlink = Hyperlink(ref=link.coordinate,
                                           location=f"'{details.title}'!A{det_row}:G{det_row}")
                link.style = "Hyperlink"

            history[row].append(status(main.cell(row, COL_RESULT).value))
        rwb.close()

    # Extra columns and summary
    header_src = main.cell(MAIN_HEADER_ROW, COL_OBJECTS_FAIL_PC)
    for col, name in ((COL_ATTEMPTS, "Attempts"), (COL_HISTORY, "Attempt\nhistory")):
        cell = main.cell(MAIN_HEADER_ROW, col, name)
        cell.font, cell.border, cell.alignment = copy(header_src.font), copy(header_src.border), copy(header_src.alignment)
    details_header = details.cell(DETAILS_HEADER_ROW, COL_DETAILS_ATTEMPT, "ATTEMPT")
    details_header.font = copy(details.cell(DETAILS_HEADER_ROW, 1).font)
    details_header.border = copy(details.cell(DETAILS_HEADER_ROW, 1).border)
    main.column_dimensions['P'].width = 22

    flaky = []
    for row, hist in history.items():
        main.cell(row, COL_ATTEMPTS, len(hist))
        main.cell(row, COL_HISTORY, " -> ".join(hist))
        if len(hist) > 1 and hist[-1] == "PASS":
            main.cell(row, COL_HISTORY).fill = FLAKY_FILL
            flaky.append(row)

    main['B5'] = 'Passed on retry:'
    main['B5'].font = copy(main['B1'].font)
    main['C5'] = len(flaky)
    main['C5'].alignment = copy(main['C1'].alignment)

    # Summary cells are formulas without cached values after openpyxl save - force recalculation
    wb.calculation.fullCalcOnLoad = True
    wb.save(output_path)
    wb.close()

    write_text_summary(main, history, flaky, output_path.with_suffix('.txt'))


def write_text_summary(main, history: dict, flaky: list, txt_path: Path):
    """Same format as XlsxReporter.save_summary_to_text_file, plus tests passed on retry"""
    results, failed_tests = [], []
    for row in main_rows(main):
        result = main.cell(row, COL_RESULT).value
        results.append(result)
        if not result:
            failed_tests.append(f"[! FAIL !] {main.cell(row, COL_SUITE).value} ({main.cell(row, COL_TEST).value})")

    total = len(results)
    passed = results.count(True)
    lines = [
        f"Tests total: {total}",
        f"Passed: {passed}",
        f"Failed: {total - passed}",
        f"Pass rate: {passed / total if total else 0:.0%}",
        "",
    ]
    lines += ["Failed tests:", *failed_tests] if failed_tests else ["Failed tests: None"]
    if flaky:
        lines += ["", "Passed on retry:"]
        lines += [f"[flaky] {main.cell(row, COL_SUITE).value} ({main.cell(row, COL_TEST).value}): "
                  f"{' -> '.join(history[row])}" for row in flaky]
    txt_path.write_text("\n".join(lines) + "\n")


def parse_args():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--base', required=True, type=Path, help='Full XLSX report from the first run')
    parser.add_argument('--retries', required=True, nargs='+', type=Path, help='XLSX reports from retries, in order')
    parser.add_argument('--output', required=True, type=Path, help='Path for merged XLSX report (.txt is written next to it)')
    return parser.parse_args()


if __name__ == '__main__':
    args = parse_args()
    merge(args.base, args.retries, args.output)
